import sys
import re
import os
import io
import tempfile
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

# Ensure stdout handles UTF-8 on Windows cp950 without crashing
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_hyperlink(paragraph, url, text, color="004B97", underline=True):
    """Add a clickable hyperlink into a python-docx paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def apply_nstc_style(doc):
    """Apply NSTC required styles: 12pt 標楷體 / Times New Roman, 1.5 line spacing."""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # 標楷體 for East Asian text
    rFonts = font.element.rPr.xpath('w:rFonts')
    if rFonts:
        rFonts[0].set(qn('w:eastAsia'), '標楷體')
    
    # 1.5 line spacing
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.5

def format_paragraph_text(p, text):
    """Parse inline markdown elements: bold (**bold**) and links [text](url)."""
    link_pattern = r'(\[.*?\]\(.*?\))'
    tokens = re.split(link_pattern, text)
    
    for token in tokens:
        link_match = re.match(r'\[(.*?)\]\((.*?)\)', token)
        if link_match:
            link_text = link_match.group(1)
            link_url = link_match.group(2)
            add_hyperlink(p, link_url, link_text)
        else:
            bold_tokens = re.split(r'(\*\*.*?\*\*)', token)
            for b_tok in bold_tokens:
                if b_tok.startswith('**') and b_tok.endswith('**'):
                    run = p.add_run(b_tok[2:-2])
                    run.bold = True
                else:
                    p.add_run(b_tok)

def load_image_source(source, base_dir):
    """Return an actual image payload and its source URL/path.

    Local Markdown paths are resolved relative to the Markdown file. HTTP(S)
    sources are downloaded before insertion. A missing or non-image source is
    a hard error: never replace an image with a text placeholder.
    """
    parsed = urlparse(source)
    if parsed.scheme in ('http', 'https'):
        request = Request(source, headers={'User-Agent': 'nstc-undergrad-grant-skill/1.0'})
        with urlopen(request, timeout=30) as response:
            content_type = response.headers.get_content_type().lower()
            data = response.read()
        if not content_type.startswith('image/'):
            raise ValueError(f'遠端來源不是圖片：{source} ({content_type})')
        if not data:
            raise ValueError(f'遠端圖片為空：{source}')
        return io.BytesIO(data), source

    local_source = source[7:] if parsed.scheme == 'file' else source
    local_path = Path(local_source)
    if not local_path.is_absolute():
        local_path = Path(base_dir) / local_path
    if not local_path.is_file():
        raise FileNotFoundError(f'找不到圖片檔：{local_path}')
    return str(local_path), str(local_path)

def parse_markdown_to_docx(md_path, docx_path):
    doc = Document()
    apply_nstc_style(doc)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    lines = content.split('\n')
    base_dir = Path(md_path).resolve().parent
    i = 0
    total_lines = len(lines)
    in_code_block = False
    
    while i < total_lines:
        line = lines[i]
        trimmed = line.strip()
        
        # Check for Code / ASCII Diagram Blocks
        if trimmed.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue
            
        if in_code_block:
            p_code = doc.add_paragraph()
            p_code.paragraph_format.line_spacing = 1.0
            p_code.paragraph_format.space_before = Pt(0)
            p_code.paragraph_format.space_after = Pt(0)
            run = p_code.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            i += 1
            continue
            
        # Check for Markdown Tables
        if trimmed.startswith('|') and trimmed.endswith('|'):
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            valid_rows = [row for row in table_lines if not re.match(r'^\|[\s\-:|]+\|$', row)]
            if valid_rows:
                num_cols = max(len([c.strip() for c in r.strip('|').split('|')]) for r in valid_rows)
                tbl = doc.add_table(rows=len(valid_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.style = 'Table Grid'
                
                for r_idx, row_str in enumerate(valid_rows):
                    cells = [c.strip() for c in row_str.strip('|').split('|')]
                    for c_idx in range(num_cols):
                        cell_text = cells[c_idx] if c_idx < len(cells) else ''
                        cell = tbl.cell(r_idx, c_idx)
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.line_spacing = 1.15
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(2)
                        format_paragraph_text(cp, cell_text)
                        if r_idx == 0:
                            for run in cp.runs:
                                run.bold = True
            continue
            
        # Check for Image Markdown: ![caption](url_or_path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', trimmed)
        if img_match:
            caption = img_match.group(1)
            img_src = img_match.group(2)
            
            p_img = doc.add_paragraph()
            p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            try:
                image_payload, resolved_source = load_image_source(img_src, base_dir)
                run_img = p_img.add_run()
                run_img.add_picture(image_payload, width=Inches(5.5))
            except Exception as e:
                raise RuntimeError(f"圖片必須是可取得的實際檔案，無法嵌入：{img_src}；原因：{e}") from e
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run_cap = p_cap.add_run(f"圖：{caption}")
            run_cap.bold = True
            
            p_src = doc.add_paragraph()
            p_src.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            p_src.add_run("（資料來源：")
            if urlparse(img_src).scheme in ('http', 'https'):
                add_hyperlink(p_src, img_src, img_src)
            else:
                p_src.add_run(str(resolved_source))
            p_src.add_run("）")
            
            i += 1
            continue
            
        # Headings
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
        elif line.startswith('## '):
            if '參考文獻' in line or 'References' in line:
                doc.add_page_break()
            p = doc.add_heading(line[3:], level=2)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith('#### '):
            p = doc.add_heading(line[5:], level=4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            format_paragraph_text(p, line[2:])
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_paragraph_text(p, line)
        elif trimmed == '' or trimmed == '---':
            pass
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            format_paragraph_text(p, line)
            
        i += 1
                    
    doc.save(docx_path)
    print(f"[OK] 成功匯出至 {docx_path}（已套用國科會標楷體 12pt、1.5 倍行高、圖表連結與出處來源格式）")

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python export_c802.py <input.md> <output.docx>")
        sys.exit(1)
        
    parse_markdown_to_docx(sys.argv[1], sys.argv[2])
